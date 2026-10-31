#!/usr/bin/env python3
"""Reads the team workbook and writes catalogue/products.json + catalogue/img/*.jpg.
   python3 -I tools/build-catalogue.py workbook.xlsx
Rule: a product is shown only if 'Show on website' is YES and 'Selling price (₹)' is a number."""
import sys, re, zipfile, json, os, io
import openpyxl
from PIL import Image
src = sys.argv[1]
# optional 2nd arg: live Google Sheet snapshot, {rows:{id:[price,show]}}; it wins over the workbook
live = json.load(open(sys.argv[2]))['rows'] if len(sys.argv) > 2 else {}
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'catalogue')
os.makedirs(out + '/img', exist_ok=True)
z = zipfile.ZipFile(src)
wb = openpyxl.load_workbook(src, data_only=True)
tabs = ['Glass Candles', 'Moulded Candles', 'Tealight Candles', 'Jesmonite Tray', 'DIY Items']
products, seen = [], {}
for i, tab in enumerate(tabs, 1):
    d = z.read(f'xl/drawings/drawing{i}.xml').decode()
    rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="\.\./media/([^"]+)"', z.read(f'xl/drawings/_rels/drawing{i}.xml.rels').decode()))
    rels.update({a: b for b, a in re.findall(r'Target="\.\./media/([^"]+)"[^>]*Id="(rId\d+)"', z.read(f'xl/drawings/_rels/drawing{i}.xml.rels').decode())})
    rowimg = {}
    for blk in re.findall(r'<xdr:(?:oneCellAnchor|twoCellAnchor|absoluteAnchor).*?</xdr:(?:oneCellAnchor|twoCellAnchor|absoluteAnchor)>', d, re.S):
        r = re.search(r'<xdr:row>(\d+)</xdr:row>', blk); e = re.search(r'r:embed="(\w+)"', blk)
        if r and e: rowimg[int(r.group(1)) + 1] = rels[e.group(1)]
    ws = wb[tab]
    hdr = [str(c.value or '').strip() for c in ws[4]]
    col = lambda name: next(j for j, h in enumerate(hdr) if h.lower().startswith(name))
    cid, cn, cp, cpr = col('product id'), col('product name'), col('product type'), col('selling price')
    cshow = next((j for j, h in enumerate(hdr) if h.lower() in ('show on website',)), None)
    if cshow is None: cshow = 11
    for rn in range(5, ws.max_row + 1):
        v = [c.value for c in ws[rn]]
        if not v[cid]: continue
        price = v[cpr]
        show = str(v[cshow] or '').strip().upper() == 'YES'
        if str(v[cid]).strip() in live: price, sh = live[str(v[cid]).strip()]; show = str(sh).upper() == 'YES'
        if not show or not isinstance(price, (int, float)): continue
        pid = str(v[cid]).strip()
        media = rowimg.get(rn)
        fn = ''
        if media:
            fn = pid.lower() + '.jpg'
            im = Image.open(io.BytesIO(z.read('xl/media/' + media))).convert('RGB')
            im.thumbnail((720, 720)); im.save(f'{out}/img/{fn}', quality=82, optimize=True)
        products.append({'id': pid, 'name': str(v[cn]).strip(), 'type': str(v[cp]).strip(), 'price': int(price) if price == int(price) else price, 'img': fn, 'sheetTab': tab, 'sheetRow': rn})
json.dump({'currency': 'INR', 'source': 'Team Master Source of Truth workbook (photos are temporary)', 'products': products}, open(out + '/products.json', 'w'), ensure_ascii=False, indent=1)
print(len(products), 'products')

# ---- static page: no JS, no cart, no login, no redirect ----
from html import escape
groups = [('Glass Candles', 'Glass candles'), ('Moulded Candles', 'Moulded candles'), ('Tealight Candles', 'Tealights'), ('Jesmonite Tray', 'Jesmonite trays & dishes'), ('DIY Items', 'DIY kits')]
sec = ''
for tab, title in groups:
    ps = [p for p in products if p['sheetTab'] == tab]
    if not ps: continue
    cards = ''.join(f'<li class="card"><img src="img/{p["img"]}" alt="{escape(p["name"])}" loading="lazy" width="360" height="360"><h3>{escape(p["name"])}</h3><p class="price">₹{p["price"]:,}</p></li>' for p in ps)
    sec += f'<section><h2>{title}</h2><ul class="grid">{cards}</ul></section>'
page = f'''<!doctype html>
<html lang="en-IN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>BUD &amp; THINGS · Catalogue</title>
<meta name="description" content="Handmade candles and Jesmonite pieces by Ashi. Catalogue with prices in rupees.">
<style>
@font-face{{font-family:'Playfair Display';src:url(../prototype/assets/fonts/playfair.woff2) format('woff2');font-weight:400 900;font-display:swap}}
@font-face{{font-family:'Montserrat';src:url(../prototype/assets/fonts/montserrat.woff2) format('woff2');font-weight:100 900;font-display:swap}}
:root{{--maroon:#8F1834;--maroon-d:#7A1229;--gold:#D6A34A;--cream:#F8F0E4;--ink:#150206}}
*{{box-sizing:border-box}}
body{{margin:0;background:var(--cream);color:var(--ink);font:400 16px/1.5 Montserrat,system-ui,sans-serif}}
header{{background:var(--maroon-d);border-bottom:3px solid var(--gold);padding:18px 16px;text-align:center}}
header img{{height:64px;width:auto;max-width:100%}}
.intro{{text-align:center;padding:32px 16px 8px}}
.intro h1{{font:600 clamp(28px,5vw,44px)/1.15 'Playfair Display',serif;color:var(--maroon);margin:0 0 8px}}
.intro p{{margin:0;color:var(--maroon-d)}}
main{{max-width:1100px;margin:0 auto;padding:0 16px 40px}}
h2{{font:600 26px 'Playfair Display',serif;color:var(--maroon);border-bottom:1px solid var(--gold);padding-bottom:6px;margin:36px 0 16px}}
.grid{{list-style:none;margin:0;padding:0;display:grid;grid-template-columns:repeat(auto-fill,minmax(210px,1fr));gap:20px}}
.card{{background:#fff;border:1px solid #e6d3b3;border-radius:6px;overflow:hidden;text-align:center;padding-bottom:14px}}
.card img{{display:block;width:100%;height:auto;aspect-ratio:1;object-fit:cover;background:var(--cream)}}
.card h3{{font:600 18px 'Playfair Display',serif;margin:12px 8px 2px}}
.price{{margin:0;font-weight:600;color:var(--maroon)}}
footer{{background:var(--maroon-d);color:var(--cream);text-align:center;padding:20px 16px;font-size:14px}}
footer b{{color:var(--gold)}}
</style></head><body>
<header><img src="../prototype/assets/silk/logo-horizontal.webp" alt="BUD &amp; THINGS by Ashi" width="280" height="64"></header>
<div class="intro"><h1>Our catalogue</h1><p>Handmade candles and Jesmonite pieces. Prices in rupees.</p></div>
<main>{sec}</main>
<footer><b>BUD &amp; THINGS</b> by Ashi · Browse only, ordering opens soon. Photos are temporary.</footer>
</body></html>
'''
open(out + '/index.html', 'w').write(page)
